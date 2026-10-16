# /// script
# dependencies = ["python-pptx==1.0.2"]
# ///
import sys
from pptx import Presentation
for path in sys.argv[1:]:
    prs = Presentation(path)
    print(f"##### {path}")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n--- slide {i}")
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t: print(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    print(" | ".join(c.text for c in row.cells))
            if sh.shape_type == 13: print("[picture]")
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n: print(f"[notes] {n}")
